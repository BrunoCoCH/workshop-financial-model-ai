"""Génère growth-curve.xlsx — atelier courbes 5 ans (skill platinn-romandee)."""
from pathlib import Path

from openpyxl import Workbook
from openpyxl.chart import LineChart, Reference
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

wb = Workbook()

navy = "1B3A4B"
teal = "2A6F7F"
light = "F4F7F8"
amber = "FFF3E0"
yellow = "FFFDE7"
total_fill = PatternFill("solid", fgColor="E3EEF1")
thin = Border(
    left=Side(style="thin", color="CFD8DC"),
    right=Side(style="thin", color="CFD8DC"),
    top=Side(style="thin", color="CFD8DC"),
    bottom=Side(style="thin", color="CFD8DC"),
)
hdr_font = Font(name="Calibri", bold=True, color="FFFFFF", size=11)
title_font = Font(name="Calibri", bold=True, color=navy, size=16)
section_font = Font(name="Calibri", bold=True, color=navy, size=12)
body = Font(name="Calibri", size=11)
hdr_fill = PatternFill("solid", fgColor=navy)
alt_fill = PatternFill("solid", fgColor=light)
note_fill = PatternFill("solid", fgColor=amber)
empty_fill = PatternFill("solid", fgColor=yellow)


def style_header_row(ws, row, headers):
    for col, h in enumerate(headers, 1):
        cell = ws.cell(row, col, h)
        cell.font = hdr_font
        cell.fill = hdr_fill
        cell.alignment = Alignment(horizontal="center", wrap_text=True)


def border_range(ws, r1, r2, c1, c2):
    for r in range(r1, r2 + 1):
        for c in range(c1, c2 + 1):
            ws.cell(r, c).border = thin


# --- Consignes ---
ws = wb.active
ws.title = "Consignes"
ws["A1"] = "Atelier — Courbes de croissance (skill financial-model-platinn-romandee)"
ws["A1"].font = title_font
ws.merge_cells("A1:D1")
ws["A3"] = "Objectif"
ws["A3"].font = section_font
ws["A4"] = (
    "Préparer les hypothèses année 1 → année 5 d’une société de services suisse, "
    "au même format de drivers que la skill financial-model-platinn-romandee "
    "(prix, volume, capacité, marge & cash). "
    "Y1 et Y5 sont donnés ; vous (ou Claude) remplissez Y2–Y4."
)
ws["A4"].alignment = Alignment(wrap_text=True)
ws.merge_cells("A4:D6")

ws["A8"] = "Déroulement"
ws["A8"].font = section_font
for i, line in enumerate(
    [
        "1. Feuille Courbes — compléter les cellules jaunes (Y2, Y3, Y4).",
        "2. Interpolation linéaire par défaut ; une rupture est OK si vous la justifiez en commentaire.",
        "3. Feuille Recap — vérifier que la capacité (FTE × heures × util.) suit la demande.",
        "4. Envoyer le tableau à la skill platinn-romandee pour générer le plan Excel complet.",
        "5. Ne pas ouvrir Corrige_formateur avant d’avoir tenté l’exercice.",
    ],
    start=9,
):
    ws.cell(i, 1, line).font = body

ws["A15"] = "Cas"
ws["A15"].font = section_font
ws["A16"] = (
    "Alpes Conseil Sàrl — conseil en organisation, Vaud, mode services (pas de SaaS, pas de VC). "
    "Devise CHF. Démarrage plan : janv. 2027. Fichier pédagogique — pas de conseil fiscal."
)
ws["A16"].font = body
ws["A16"].alignment = Alignment(wrap_text=True)
ws.merge_cells("A16:D17")

ws["A19"] = "Mapping skill (rappel)"
ws["A19"].font = section_font
ws["A20"] = (
    "NewClients, ClientChurn, ProjectDemand, TargetUtil, RetainerFee/BillRate (via PriceIncrease), "
    "roles / FTE delivery → inputs JSON de financial-model-platinn-romandee (voir sample lemanconseil)."
)
ws["A20"].font = body
ws["A20"].alignment = Alignment(wrap_text=True)
ws.merge_cells("A20:D21")
ws.column_dimensions["A"].width = 100

# --- Courbes ---
# Structure: (driver_key, label, unit, y1, y5, format)
# Y2-Y4 left empty for exercise
DRIVERS = [
    # Driver 1 Prix
    ("RetainerFee", "1. Prix — Honoraire retainer mensuel type", "CHF/mois", 3800, 4500, "chf"),
    ("BillRate", "1. Prix — Taux horaire projet", "CHF/h", 175, 195, "chf"),
    ("PriceIncrease", "1. Prix — Hausse prix annuelle (à partir de Y2)", "%", None, 0.03, "pct"),
    # Driver 2 Volume
    ("StartClients", "2. Volume — Clients retainer en début de plan (fixe)", "nb", 5, 5, "num"),
    ("NewClients", "2. Volume — Nouveaux clients retainer / mois", "nb/mois", 0.5, 1.2, "num"),
    ("ClientChurn", "2. Volume — Churn clients / mois", "%", 0.025, 0.018, "pct"),
    ("RetainerHours", "2. Volume — Heures / client retainer / mois", "h", 20, 22, "num"),
    ("ProjectDemand", "2. Volume — Demande projets (heures/mois)", "h/mois", 120, 400, "num"),
    # Driver 3 Capacité
    ("DeliveryFTE", "3. Capacité — FTE Delivery (moyenne année)", "FTE", 3.0, 7.0, "num"),
    ("HoursPerFTE", "3. Capacité — Heures travaillées / FTE / mois", "h", 150, 150, "num"),
    ("TargetUtil", "3. Capacité — Utilisation cible", "%", 0.70, 0.76, "pct"),
    ("SubShare", "3. Capacité — Part overflow en freelance", "%", 0.4, 0.5, "pct"),
    ("SubCost", "3. Capacité — Coût horaire freelance", "CHF/h", 115, 130, "chf"),
    # Driver 4 Cash / structure
    ("Rent", "4. Cash — Loyer mensuel", "CHF/mois", 3200, 5200, "chf"),
    ("DSO", "4. Cash — Délai encaissement clients", "jours", 45, 40, "num"),
    ("MinCash", "4. Cash — Plancher de trésorerie", "CHF", 40000, 80000, "chf"),
]

c = wb.create_sheet("Courbes")
c["A1"] = "Alpes Conseil Sàrl — Courbes d’hypothèses (5 ans)"
c["A1"].font = title_font
c.merge_cells("A1:H1")
c["A2"] = (
    "Jaune = à compléter (Y2–Y4). Y1 et Y5 fournis. "
    "Aligné drivers skill financial-model-platinn-romandee (mode services)."
)
c["A2"].font = Font(name="Calibri", italic=True, size=11, color=teal)
c.merge_cells("A2:H2")

headers = ["Clé skill", "Hypothèse", "Unité", "Y1", "Y2", "Y3", "Y4", "Y5"]
style_header_row(c, 4, headers)

for i, (key, label, unit, y1, y5, fmt) in enumerate(DRIVERS):
    r = 5 + i
    c.cell(r, 1, key).font = Font(name="Calibri", size=10, color=teal)
    c.cell(r, 2, label).font = body
    c.cell(r, 3, unit).font = body
    # Y1
    if key == "PriceIncrease":
        c.cell(r, 4, "—")  # hausse from Y2
        c.cell(r, 4).alignment = Alignment(horizontal="center")
    else:
        c.cell(r, 4, y1)
    # Y2-Y4 empty (except HoursPerFTE flat hint: leave empty still for consistency)
    for col in (5, 6, 7):
        c.cell(r, col, None)
        c.cell(r, col).fill = empty_fill
    # Y5
    c.cell(r, 8, y5)

    for col in range(1, 9):
        c.cell(r, col).border = thin
        if i % 2 and col not in (5, 6, 7):
            if c.cell(r, col).fill.fgColor is None or str(
                getattr(c.cell(r, col).fill.fgColor, "rgb", "")
            ) in ("00000000", "None"):
                c.cell(r, col).fill = alt_fill

    # number formats
    for col in (4, 5, 6, 7, 8):
        cell = c.cell(r, col)
        if cell.value == "—" or cell.value is None:
            continue
        if fmt == "chf":
            cell.number_format = "#,##0"
        elif fmt == "pct":
            cell.number_format = "0.0%"
        else:
            cell.number_format = "0.0" if isinstance(cell.value, float) else "0"

# Special: PriceIncrease Y2-Y4 yellow, Y5 given; Y1 is em-dash
# Already set Y1 as —. Put hint that Y2 can start at 2%
c["A22"] = "Indices"
c["A22"].font = section_font
c["A23"] = (
    "Interpolation linéaire : valeur_Yn = Y1 + (Y5−Y1)×(n−1)/4. "
    "Pour PriceIncrease, Y1 n’existe pas (hausse dès Y2) : proposez une montée vers 3 %. "
    "StartClients reste constant (stock initial, pas une courbe)."
)
c["A23"].fill = note_fill
c["A23"].alignment = Alignment(wrap_text=True)
c.merge_cells("A23:H24")

widths = [16, 48, 12, 10, 10, 10, 10, 10]
for i, w in enumerate(widths, 1):
    c.column_dimensions[get_column_letter(i)].width = w

# --- Recap (formulas once Y2-Y4 filled; capacity check uses Y1 and Y5 only + placeholders) ---
rcap = wb.create_sheet("Recap")
rcap["A1"] = "Recap capacité vs demande (après remplissage des Courbes)"
rcap["A1"].font = title_font
rcap.merge_cells("A1:F1")
rcap["A2"] = (
    "Quand Y2–Y4 sont remplis, ces indicateurs se calculent depuis Courbes. "
    "Capacité billable ≈ DeliveryFTE × HoursPerFTE × TargetUtil."
)
rcap["A2"].font = Font(name="Calibri", italic=True, size=11, color=teal)
rcap.merge_cells("A2:F2")

style_header_row(
    rcap, 4, ["Indicateur", "Y1", "Y2", "Y3", "Y4", "Y5"]
)

# Row refs on Courbes: NewClients=9, ProjectDemand=12, DeliveryFTE=13, HoursPerFTE=14, TargetUtil=15
# RetainerHours=11, StartClients=8
indicators = [
    ("Nouveaux clients / mois", 9),
    ("Demande projets h/mois", 12),
    ("FTE Delivery", 13),
    ("Utilisation cible", 15),
    ("Capacité billable h/mois", None),  # formula
    ("Heures retainer (approx clients×h)*", None),
]

# Simple formulas linking to Courbes columns D-H (Y1-Y5) = cols 4-8
rcap["A5"] = "Nouveaux clients / mois"
rcap["A6"] = "Demande projets (h/mois)"
rcap["A7"] = "FTE Delivery"
rcap["A8"] = "Utilisation cible"
rcap["A9"] = "Capacité billable (h/mois)"
rcap["A10"] = "Demande totale approx. (projets + retainer start×h)"

for col_i, col in enumerate("DEFGH", start=2):
    # NewClients row 9 on Courbes
    rcap.cell(5, col_i, f"=Courbes!{col}9")
    rcap.cell(5, col_i).number_format = "0.00"
    rcap.cell(6, col_i, f"=Courbes!{col}12")
    rcap.cell(6, col_i).number_format = "#,##0"
    rcap.cell(7, col_i, f"=Courbes!{col}13")
    rcap.cell(7, col_i).number_format = "0.0"
    rcap.cell(8, col_i, f"=Courbes!{col}15")
    rcap.cell(8, col_i).number_format = "0.0%"
    # Capacity = FTE * Hours * Util  → Courbes rows 13, 14, 15
    rcap.cell(9, col_i, f"=Courbes!{col}13*Courbes!{col}14*Courbes!{col}15")
    rcap.cell(9, col_i).number_format = "#,##0"
    # Demand approx = ProjectDemand + StartClients * RetainerHours
    rcap.cell(10, col_i, f"=Courbes!{col}12+Courbes!{col}8*Courbes!{col}11")
    rcap.cell(10, col_i).number_format = "#,##0"

for r in range(5, 11):
    rcap.cell(r, 1).font = body
    for col in range(1, 7):
        rcap.cell(r, col).border = thin
        if r in (9, 10):
            rcap.cell(r, col).fill = total_fill

rcap["A12"] = "Lecture"
rcap["A12"].font = section_font
rcap["A13"] = (
    "Si Demande totale ≫ Capacité billable sur une année, il faut plus de FTE Delivery "
    "(ou moins de ProjectDemand / plus de freelances). C’est le cœur du driver 3 de la skill."
)
rcap["A13"].fill = note_fill
rcap["A13"].alignment = Alignment(wrap_text=True)
rcap.merge_cells("A13:F14")

rcap["A16"] = "* Approx. pédagogique : stock initial × heures ; la skill gère le stock de clients dynamiquement."
rcap["A16"].font = Font(name="Calibri", italic=True, size=9, color="666666")

for i, w in enumerate([48, 12, 12, 12, 12, 12], 1):
    rcap.column_dimensions[get_column_letter(i)].width = w

# Chart on Recap for ProjectDemand once filled — chart from Courbes Y1/Y5 only won't show mid; skip or chart capacity after fill
chart = LineChart()
chart.title = "Demande projets vs capacité (après remplissage)"
chart.style = 10
chart.y_axis.title = "Heures / mois"
chart.x_axis.title = "Année"
data = Reference(rcap, min_col=2, min_row=6, max_col=6, max_row=9)
cats = Reference(rcap, min_col=2, min_row=4, max_col=6, max_row=4)
# Better categories: use manual labels - skip cats from header; create year labels row
rcap["B3"] = "Y1"
rcap["C3"] = "Y2"
rcap["D3"] = "Y3"
rcap["E3"] = "Y4"
rcap["F3"] = "Y5"
cats = Reference(rcap, min_col=2, min_row=3, max_col=6, max_row=3)
data = Reference(rcap, min_col=1, min_row=6, max_col=6, max_row=9)
chart.add_data(data, from_rows=True, titles_from_data=True)
chart.set_categories(cats)
chart.shape = 4
chart.width = 15
chart.height = 8
rcap.add_chart(chart, "A18")

# --- Prompt skill ---
p = wb.create_sheet("Prompt_skill")
p["A1"] = "Prompt type — enchaîner avec financial-model-platinn-romandee"
p["A1"].font = title_font
p.merge_cells("A1:B1")
p["A3"] = (
    "J’ai préparé les courbes 5 ans d’Alpes Conseil Sàrl (mode services, pas de subscriptions, "
    "pas de VC) dans growth-curve.xlsx, feuille Courbes. "
    "Utilise la skill financial-model-platinn-romandee : parte de ces hypothèses, "
    "pose-moi seulement les questions manquantes (équipe détaillée, canton, BFR, etc.), "
    "puis génère le JSON et le plan Excel."
)
p["A3"].alignment = Alignment(wrap_text=True)
p.merge_cells("A3:B6")
p["A8"] = "Variante sans la skill (Claude in Excel / Cursor)"
p["A8"].font = section_font
p["A9"] = (
    "Sur la feuille Courbes, remplis Y2 Y3 Y4 par interpolation linéaire entre Y1 et Y5 "
    "(sauf StartClients constant et PriceIncrease dès Y2). "
    "Écris les formules Excel, pas des valeurs en dur si possible."
)
p["A9"].alignment = Alignment(wrap_text=True)
p.merge_cells("A9:B11")
p.column_dimensions["A"].width = 100

# --- Corrigé ---
corr = wb.create_sheet("Corrige_formateur")
corr["A1"] = "Corrigé indicatif — interpolation linéaire"
corr["A1"].font = title_font
corr["A3"] = "Formule"
corr["A3"].font = section_font
corr["A4"] = "Yn = Y1 + (Y5 − Y1) × (n − 1) / 4    pour n = 2,3,4"
corr["A6"] = "Exemples"
corr["A6"].font = section_font
examples = [
    "NewClients : 0.5 → 0.675 → 0.85 → 1.025 → 1.2",
    "ProjectDemand : 120 → 190 → 260 → 330 → 400",
    "DeliveryFTE : 3.0 → 4.0 → 5.0 → 6.0 → 7.0",
    "TargetUtil : 70 % → 71.5 % → 73 % → 74.5 % → 76 %",
    "BillRate : 175 → 180 → 185 → 190 → 195",
    "PriceIncrease : — → 2.0 % → 2.5 % → 2.75 % → 3.0 % (exemple de montée)",
    "StartClients : 5 · 5 · 5 · 5 · 5",
    "HoursPerFTE : 150 constant",
]
for i, e in enumerate(examples, start=7):
    corr.cell(i, 1, e).font = body

corr["A16"] = "Point pédagogique capacité"
corr["A16"].font = section_font
corr["A17"] = (
    "Avec interpolation linéaire FTE et demande, Y5 : capacité ≈ 7×150×0.76 = 798 h ; "
    "demande approx. 400 + 5×22 = 510 h → marge de capacité OK. "
    "Si les participants mettent la demande trop haut sans monter les FTE, le Recap le montre."
)
corr["A17"].alignment = Alignment(wrap_text=True)
corr.merge_cells("A17:B19")
corr["A21"] = "Ne montrez pas cette feuille avant la fin de l’exercice."
corr["A21"].fill = note_fill
corr["A21"].font = Font(name="Calibri", bold=True, size=11)
corr.column_dimensions["A"].width = 95

out = Path(__file__).resolve().parent / "growth-curve.xlsx"
wb.save(out)
print(out)
