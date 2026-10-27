"""Génère le classeur d'exercice Claude dans Excel."""
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

wb = Workbook()

navy = "1B3A4B"
teal = "2A6F7F"
light = "F4F7F8"
amber = "FFF3E0"
green = "E8F5E9"
thin = Border(
    left=Side(style="thin", color="CFD8DC"),
    right=Side(style="thin", color="CFD8DC"),
    top=Side(style="thin", color="CFD8DC"),
    bottom=Side(style="thin", color="CFD8DC"),
)
hdr_font = Font(name="Calibri", bold=True, color="FFFFFF", size=11)
title_font = Font(name="Calibri", bold=True, color=navy, size=16)
section_font = Font(name="Calibri", bold=True, color=navy, size=12)
body = Font(name="Calibri", size=11)
hdr_fill = PatternFill("solid", fgColor=navy)
alt_fill = PatternFill("solid", fgColor=light)
note_fill = PatternFill("solid", fgColor=amber)
todo_fill = PatternFill("solid", fgColor="FFF8E1")
ok_fill = PatternFill("solid", fgColor=green)


def chf(cell):
    cell.number_format = "#,##0"
    cell.font = body
    cell.alignment = Alignment(horizontal="right")


def pct(cell):
    cell.number_format = "0.0%"
    cell.font = body


# ========== Consignes ==========
ws = wb.active
ws.title = "Consignes"
ws["A1"] = "Atelier — Plugin Claude dans Excel"
ws["A1"].font = title_font
ws.merge_cells("A1:C1")
ws["A3"] = "Objectif"
ws["A3"].font = section_font
ws["A4"] = (
    "Utiliser l’add-in Claude for Excel pour compléter un tableau de pilotage mensuel : "
    "formules d’écarts, totaux, commentaire de gestion, puis un petit correctif. "
    "Tout se passe dans Excel — pas besoin de Cursor pour cet atelier."
)
ws["A4"].alignment = Alignment(wrap_text=True)
ws.merge_cells("A4:C6")

ws["A8"] = "Installation rapide"
ws["A8"].font = section_font
for i, line in enumerate(
    [
        "1. Excel Microsoft 365 → Insertion (ou Accueil) → Compléments → chercher « Claude ».",
        "2. Connectez votre compte Anthropic si demandé.",
        "3. Ouvrez le panneau Claude à droite ; gardez ce classeur actif.",
        "4. Pour chaque mission : sélectionnez la plage utile si possible, puis collez le prompt.",
    ],
    start=9,
):
    ws.cell(i, 1, line).font = body

ws["A14"] = "Règles d’or (à rappeler aux participants)"
ws["A14"].font = section_font
for i, line in enumerate(
    [
        "• Demandez des FORMULES Excel, pas des chiffres copiés en dur.",
        "• Une mission à la fois — vérifiez le résultat avant la suivante.",
        "• Si Claude écrit à côté : précisez la plage exacte (ex. D5:D12).",
        "• Relisez toujours : Claude peut se tromper de signe ou de total.",
    ],
    start=15,
):
    ws.cell(i, 1, line).font = body

ws["A20"] = "Société fictive"
ws["A20"].font = section_font
ws["A21"] = (
    "Alpes Conseil Sàrl — conseil, Vaud. Devise CHF. Fichier pédagogique uniquement."
)
ws["A21"].font = body
ws.column_dimensions["A"].width = 100

# ========== Missions ==========
mis = wb.create_sheet("Missions")
mis["A1"] = "4 missions — à faire dans l’ordre avec Claude in Excel"
mis["A1"].font = title_font
mis.merge_cells("A1:B1")

missions = [
    (
        "Mission 1 — Formules d’écart (5 min)",
        "Feuille Donnees. Les colonnes Budget et Realise sont remplies ; "
        "Ecart et Ecart_pct sont vides.\n\n"
        "Prompt à coller dans Claude :\n"
        "« Sur la feuille Donnees, pour les postes des lignes 5 à 12 : "
        "dans la colonne Ecart (D) écris des formules =Realise−Budget ; "
        "dans Ecart_pct (E) des formules d’écart en pourcentage du budget. "
        "Ne modifie pas les colonnes A–C. Utilise des formules, pas des valeurs. »",
        "Contrôle : D5 doit être une formule (ex. =C5-B5), pas −18000 en dur.",
    ),
    (
        "Mission 2 — Totaux (5 min)",
        "Toujours feuille Donnees. Les lignes Total produits / Total charges / Resultat "
        "sont préparées mais sans formules.\n\n"
        "Prompt :\n"
        "« Complète B14:E16 avec des formules : "
        "Total produits = honoraires + autres produits ; "
        "Total charges = somme des 6 lignes de charges ; "
        "Resultat = produits − charges. "
        "Pour les colonnes Ecart et Ecart_pct des totaux, utilise aussi des formules. »",
        "Contrôle : B16 ≈ 27'000 en budget ; le réalisé doit être plus bas.",
    ),
    (
        "Mission 3 — Commentaire de gestion (5 min)",
        "Feuille Synthese : zone Commentaire_Claude (B3) est vide.\n\n"
        "Prompt :\n"
        "« Lis le tableau Donnees (y compris les totaux). "
        "Écris dans Synthese!B3 un commentaire de gestion en français, 4–6 phrases max : "
        "ce qui tire le résultat, 2 causes probables à confirmer, 1 action pour avril. "
        "CHF avec apostrophes. N’invente pas de chiffres absents du tableau. »",
        "Contrôle : le texte cite le CA et idéalement les freelances ; une action claire.",
    ),
    (
        "Mission 4 — Trouver l’erreur (5 min)",
        "Feuille Piege : un écart « favorable » est mal signé / mal libellé.\n\n"
        "Prompt :\n"
        "« Sur la feuille Piege, trouve l’erreur de logique dans le calcul ou le libellé "
        "des écarts. Corrige uniquement ce qui est faux (formule ou texte Sens). "
        "Explique en une phrase dans B12. »",
        "Contrôle : voir Corrige_formateur — le CA en baisse ne peut pas être « favorable ».",
    ),
]

row = 3
for title, body_txt, check in missions:
    mis.cell(row, 1, title).font = section_font
    mis.cell(row, 1).fill = todo_fill
    mis.merge_cells(start_row=row, start_column=1, end_row=row, end_column=2)
    row += 1
    mis.cell(row, 1, body_txt).font = body
    mis.cell(row, 1).alignment = Alignment(wrap_text=True, vertical="top")
    mis.merge_cells(start_row=row, start_column=1, end_row=row + 5, end_column=2)
    mis.row_dimensions[row].height = 90
    row += 6
    mis.cell(row, 1, "✓ " + check).font = Font(name="Calibri", italic=True, size=10, color=teal)
    mis.merge_cells(start_row=row, start_column=1, end_row=row, end_column=2)
    row += 2

mis.column_dimensions["A"].width = 100
mis.column_dimensions["B"].width = 20

# ========== Donnees ==========
d = wb.create_sheet("Donnees")
d["A1"] = "Alpes Conseil Sàrl — Pilotage mensuel"
d["A1"].font = title_font
d.merge_cells("A1:E1")
d["A2"] = "Période : mars 2026  |  À compléter avec Claude (missions 1 et 2)"
d["A2"].font = Font(name="Calibri", italic=True, size=11, color=teal)

headers = ["Poste", "Budget", "Realise", "Ecart", "Ecart_pct"]
for col, h in enumerate(headers, 1):
    cell = d.cell(4, col, h)
    cell.font = hdr_font
    cell.fill = hdr_fill

data = [
    ("CA_honoraires", 100_000, 82_000),
    ("Autres_produits", 2_000, 1_500),
    ("Personnel", 55_000, 56_500),
    ("Freelances", 8_000, 12_000),
    ("Loyer", 4_500, 4_500),
    ("Marketing", 3_000, 4_200),
    ("Informatique", 1_800, 1_750),
    ("Autres_charges", 2_700, 2_400),
]
for i, (poste, bud, real) in enumerate(data):
    r = 5 + i
    d.cell(r, 1, poste).font = body
    d.cell(r, 2, bud)
    chf(d.cell(r, 2))
    d.cell(r, 3, real)
    chf(d.cell(r, 3))
    # Ecart / Ecart_pct intentionally empty for Mission 1
    d.cell(r, 4, None)
    d.cell(r, 5, None)
    for col in range(1, 6):
        d.cell(r, col).border = thin
        if i % 2:
            d.cell(r, col).fill = alt_fill
    # Hint fill on empty formula cells
    d.cell(r, 4).fill = PatternFill("solid", fgColor="FFFDE7")
    d.cell(r, 5).fill = PatternFill("solid", fgColor="FFFDE7")

d["A14"] = "Total_produits"
d["A15"] = "Total_charges"
d["A16"] = "Resultat"
for r, label in ((14, "Total_produits"), (15, "Total_charges"), (16, "Resultat")):
    d.cell(r, 1, label).font = Font(name="Calibri", bold=True, size=11)
    for col in range(1, 6):
        d.cell(r, col).border = thin
        d.cell(r, col).fill = PatternFill("solid", fgColor="E3EEF1")
    # leave B–E empty for Mission 2
    for col in (2, 3, 4, 5):
        d.cell(r, col).fill = PatternFill("solid", fgColor="FFFDE7")

d["A18"] = "Zones jaunes = à faire remplir par Claude (formules)."
d["A18"].fill = note_fill
d["A18"].font = Font(name="Calibri", italic=True, size=10)
d.merge_cells("A18:E18")

for i, w in enumerate([22, 12, 12, 12, 12], 1):
    d.column_dimensions[get_column_letter(i)].width = w

# ========== Synthese ==========
s = wb.create_sheet("Synthese")
s["A1"] = "Synthèse COMEX — mars 2026"
s["A1"].font = title_font
s["A3"] = "Commentaire_Claude"
s["A3"].font = section_font
s["B3"] = ""  # Mission 3 writes here
s["B3"].fill = PatternFill("solid", fgColor="FFFDE7")
s["B3"].alignment = Alignment(wrap_text=True, vertical="top")
s.merge_cells("B3:D10")
s["A3"].alignment = Alignment(vertical="top")
s["A12"] = "Après Mission 3, le commentaire doit apparaître dans la zone jaune B3."
s["A12"].font = Font(name="Calibri", italic=True, size=10, color=teal)
s.column_dimensions["A"].width = 22
s.column_dimensions["B"].width = 28
s.column_dimensions["C"].width = 28
s.column_dimensions["D"].width = 28
s.row_dimensions[3].height = 20

# ========== Piege ==========
p = wb.create_sheet("Piege")
p["A1"] = "Mini-cas — trouvez l’erreur (Mission 4)"
p["A1"].font = title_font
p["A2"] = "Un stagiaire a « aidé » sur les écarts. Une ligne est incohérente."
p["A2"].font = Font(name="Calibri", italic=True, size=11, color=teal)

for col, h in enumerate(["Poste", "Budget", "Realise", "Ecart", "Sens"], 1):
    cell = p.cell(4, col, h)
    cell.font = hdr_font
    cell.fill = hdr_fill

# Intentional error: CA down but Sens = Favorable; Ecart formula wrong sign (Budget-Realise)
trap = [
    ("CA_honoraires", 100_000, 82_000, "=B5-C5", "Favorable"),  # WRONG sens + inverted formula
    ("Freelances", 8_000, 12_000, "=C6-B6", "Défavorable"),  # OK if Ecart = Realise-Budget for cost... 
    # For costs, higher = unfavorable. C6-B6 = +4000 défavorable — OK
]
# Simplify: only one clear error on CA
p["A5"] = "CA_honoraires"
p["B5"] = 100_000
chf(p["B5"])
p["C5"] = 82_000
chf(p["C5"])
p["D5"] = "=B5-C5"  # inverted vs convention Realise-Budget used elsewhere
chf(p["D5"])
p["E5"] = "Favorable"  # wrong label

p["A6"] = "Freelances"
p["B6"] = 8_000
chf(p["B6"])
p["C6"] = 12_000
chf(p["C6"])
p["D6"] = "=C6-B6"
chf(p["D6"])
p["E6"] = "Défavorable"

p["A7"] = "Marketing"
p["B7"] = 3_000
chf(p["B7"])
p["C7"] = 4_200
chf(p["C7"])
p["D7"] = "=C7-B7"
chf(p["D7"])
p["E7"] = "Défavorable"

for r in range(5, 8):
    for col in range(1, 6):
        p.cell(r, col).border = thin

p["A10"] = "Convention du classeur : Ecart = Realise − Budget. Pour le CA, un écart négatif = défavorable."
p["A10"].font = Font(name="Calibri", size=10, italic=True)
p.merge_cells("A10:E10")

p["A12"] = "Explication (Claude) :"
p["A12"].font = section_font
p["B12"] = ""
p["B12"].fill = PatternFill("solid", fgColor="FFFDE7")
p.merge_cells("B12:E12")

for i, w in enumerate([18, 12, 12, 12, 14], 1):
    p.column_dimensions[get_column_letter(i)].width = w

# ========== Corrigé ==========
c = wb.create_sheet("Corrige_formateur")
c["A1"] = "Corrigé indicatif (formateur)"
c["A1"].font = title_font
c["A3"] = "Mission 1"
c["A3"].font = section_font
c["A4"] = "D5: =C5-B5  →  -18000 ;  E5: =D5/B5  →  -18,0 %  (et relative pour les lignes 5–12)."
c["A6"] = "Mission 2"
c["A6"].font = section_font
c["A7"] = "B14:=B5+B6 ; B15:=SUM(B7:B12) ; B16:=B14-B15  (idem colonne C). D/E en formules."
c["A8"] = "Budget résultat attendu : 102'000 − 75'000 = 27'000. Réalisé : 83'500 − 81'350 = 2'150."
c["A10"] = "Mission 3"
c["A10"].font = section_font
c["A11"] = (
    "Le commentaire doit mettre en avant le CA (−18k) et les freelances (+4k), "
    "mentionner des causes à confirmer (missions reportées / absence), "
    "et une action (forecast avril, plafond freelance)."
)
c["A11"].alignment = Alignment(wrap_text=True)
c.merge_cells("A11:B13")
c["A15"] = "Mission 4"
c["A15"].font = section_font
c["A16"] = (
    "Erreur : sur CA_honoraires, Ecart utilise Budget−Realise (=B5-C5) et Sens dit Favorable. "
    "Correction : D5 =C5-B5 (ou garder la valeur mais Sens = Défavorable). "
    "Un CA inférieur au budget est défavorable."
)
c["A16"].alignment = Alignment(wrap_text=True)
c.merge_cells("A16:B18")
c["A20"] = "Ne montrez pas cette feuille avant la fin de l’atelier."
c["A20"].fill = note_fill
c["A20"].font = Font(name="Calibri", bold=True, size=11)
c.column_dimensions["A"].width = 100

out = Path(__file__).resolve().parent / "Pilotage_Mensuel_AlpesConseil.xlsx"
wb.save(out)
print(out)
