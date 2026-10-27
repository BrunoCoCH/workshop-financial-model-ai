"""Génère le classeur d'exercice pour la skill budget-variance-note."""
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

wb = Workbook()

navy = "1B3A4B"
teal = "2A6F7F"
light = "F4F7F8"
amber = "FFF3E0"
thin = Border(
    left=Side(style="thin", color="CFD8DC"),
    right=Side(style="thin", color="CFD8DC"),
    top=Side(style="thin", color="CFD8DC"),
    bottom=Side(style="thin", color="CFD8DC"),
)
hdr_font = Font(name="Calibri", bold=True, color="FFFFFF", size=11)
title_font = Font(name="Calibri", bold=True, color=navy, size=16)
section_font = Font(name="Calibri", bold=True, color=navy, size=12)
body = Font(name="Calibri", size=11)
hdr_fill = PatternFill("solid", fgColor=navy)
alt_fill = PatternFill("solid", fgColor=light)
note_fill = PatternFill("solid", fgColor=amber)
total_fill = PatternFill("solid", fgColor="E3EEF1")


def chf(cell):
    cell.number_format = "#,##0"
    cell.font = body
    cell.alignment = Alignment(horizontal="right")


def pct(cell):
    cell.number_format = "0.0%"
    cell.font = body


# --- Consignes ---
ws = wb.active
ws.title = "Consignes"
ws["A1"] = "Atelier — Créer une skill : budget-variance-note"
ws["A1"].font = title_font
ws.merge_cells("A1:D1")
ws["A3"] = "Objectif"
ws["A3"].font = section_font
ws["A4"] = (
    "Utilisez la feuille « Mars » (ou « Trimestre ») avec votre skill budget-variance-note. "
    "Demandez à Claude / Cursor : « Rédige la note d'écarts à partir de ce fichier ». "
    "La skill doit produire synthèse, tableau des 3 postes clés, questions métier "
    "et une prochaine action."
)
ws["A4"].alignment = Alignment(wrap_text=True)
ws.merge_cells("A4:D6")

ws["A8"] = "Déroulement suggéré (15–20 min)"
ws["A8"].font = section_font
steps = [
    "1. Ouvrez Create-Skill et créez la skill budget-variance-note (modèle fourni en atelier).",
    "2. Ouvrez ce classeur — ne lisez que la feuille Mars pour le premier essai.",
    "3. Lancez la skill sur les chiffres de Mars (sans regarder la feuille Corrigé).",
    "4. Comparez votre note à la feuille Corrigé_formateur (ordre de grandeur, pas mot à mot).",
    "5. Variante : refaites l'exercice avec la feuille Trimestre.",
]
for i, s in enumerate(steps):
    ws.cell(9 + i, 1, s).font = body

ws["A15"] = "Société fictive"
ws["A15"].font = section_font
ws["A16"] = (
    "Alpes Conseil Sàrl — cabinet de conseil, Vaud. Devise : CHF. "
    "Fichier pédagogique uniquement — pas de conseil fiscal ou financier."
)
ws["A16"].font = body
ws.column_dimensions["A"].width = 100

# --- Mars ---
m = wb.create_sheet("Mars")
m["A1"] = "Alpes Conseil Sàrl — Budget vs réalisé"
m["A1"].font = title_font
m.merge_cells("A1:F1")
m["A2"] = "Période : mars 2026  |  Comparaison : budget mensuel vs réalisé"
m["A2"].font = Font(name="Calibri", italic=True, size=11, color=teal)

headers = [
    "Poste",
    "Budget CHF",
    "Réalisé CHF",
    "Écart CHF",
    "Écart %",
    "Commentaire métier (brut)",
]
for col, h in enumerate(headers, 1):
    cell = m.cell(4, col, h)
    cell.font = hdr_font
    cell.fill = hdr_fill
    cell.alignment = Alignment(horizontal="center")

rows = [
    ("Chiffre d'affaires — honoraires", 100_000, 82_000,
     "2 missions reportées ; 1 client a réduit le volume"),
    ("Autres produits", 2_000, 1_500, ""),
    ("Charges de personnel", 55_000, 56_500, "Heures supp. sur un livrable urgent"),
    ("Sous-traitance / freelances", 8_000, 12_000, "Remplacement d'un consultant malade"),
    ("Loyer et charges locatives", 4_500, 4_500, ""),
    ("Marketing et acquisition", 3_000, 4_200, "Campagne LinkedIn non budgétée"),
    ("Informatique / outils", 1_800, 1_750, ""),
    ("Autres frais généraux", 2_700, 2_400, ""),
]

for i, (poste, bud, real, com) in enumerate(rows):
    r = 5 + i
    m.cell(r, 1, poste).font = body
    m.cell(r, 2, bud)
    chf(m.cell(r, 2))
    m.cell(r, 3, real)
    chf(m.cell(r, 3))
    m.cell(r, 4, f"=C{r}-B{r}")
    chf(m.cell(r, 4))
    m.cell(r, 5, f"=IF(B{r}=0,0,D{r}/B{r})")
    pct(m.cell(r, 5))
    m.cell(r, 6, com).font = body
    for col in range(1, 7):
        m.cell(r, col).border = thin
        if i % 2:
            m.cell(r, col).fill = alt_fill

m["A14"] = "Total produits"
m["A15"] = "Total charges"
m["A16"] = "Résultat d'exploitation (approx.)"
m["B14"] = "=B5+B6"
m["C14"] = "=C5+C6"
m["B15"] = "=SUM(B7:B12)"
m["C15"] = "=SUM(C7:C12)"
m["B16"] = "=B14-B15"
m["C16"] = "=C14-C15"
for r in (14, 15, 16):
    m.cell(r, 4, f"=C{r}-B{r}")
    m.cell(r, 5, f"=IF(B{r}=0,0,D{r}/B{r})")
    for col in range(1, 6):
        m.cell(r, col).font = Font(name="Calibri", bold=True, size=11)
        m.cell(r, col).fill = total_fill
        m.cell(r, col).border = thin
    chf(m.cell(r, 2))
    chf(m.cell(r, 3))
    chf(m.cell(r, 4))
    pct(m.cell(r, 5))

m["A18"] = "Astuce atelier"
m["A18"].font = section_font
m["A19"] = (
    "La skill doit garder au plus 3 postes qui expliquent l'essentiel de l'écart. "
    "Les commentaires métier sont des indices bruts — la note ne doit pas les recopier tels quels."
)
m["A19"].fill = note_fill
m["A19"].alignment = Alignment(wrap_text=True)
m.merge_cells("A19:F20")

for i, w in enumerate([36, 14, 14, 12, 10, 52], 1):
    m.column_dimensions[get_column_letter(i)].width = w

# --- Trimestre ---
t = wb.create_sheet("Trimestre")
t["A1"] = "Alpes Conseil Sàrl — Budget vs réalisé Q1 2026"
t["A1"].font = title_font
t.merge_cells("A1:E1")
t["A2"] = "Période : janvier–mars 2026  |  Variante pour un 2e essai"
t["A2"].font = Font(name="Calibri", italic=True, size=11, color=teal)

for col, h in enumerate(
    ["Poste", "Budget CHF", "Réalisé CHF", "Écart CHF", "Écart %"], 1
):
    cell = t.cell(4, col, h)
    cell.font = hdr_font
    cell.fill = hdr_fill

q_rows = [
    ("Chiffre d'affaires — honoraires", 280_000, 251_000),
    ("Autres produits", 6_000, 5_200),
    ("Charges de personnel", 165_000, 168_000),
    ("Sous-traitance / freelances", 24_000, 31_000),
    ("Loyer et charges locatives", 13_500, 13_500),
    ("Marketing et acquisition", 9_000, 11_500),
    ("Informatique / outils", 5_400, 5_100),
    ("Autres frais généraux", 8_100, 7_800),
]
for i, (poste, bud, real) in enumerate(q_rows):
    r = 5 + i
    t.cell(r, 1, poste).font = body
    t.cell(r, 2, bud)
    chf(t.cell(r, 2))
    t.cell(r, 3, real)
    chf(t.cell(r, 3))
    t.cell(r, 4, f"=C{r}-B{r}")
    chf(t.cell(r, 4))
    t.cell(r, 5, f"=IF(B{r}=0,0,D{r}/B{r})")
    pct(t.cell(r, 5))
    for col in range(1, 6):
        t.cell(r, col).border = thin
        if i % 2:
            t.cell(r, col).fill = alt_fill

t["A14"] = "Total produits"
t["A15"] = "Total charges"
t["A16"] = "Résultat d'exploitation (approx.)"
t["B14"] = "=B5+B6"
t["C14"] = "=C5+C6"
t["B15"] = "=SUM(B7:B12)"
t["C15"] = "=SUM(C7:C12)"
t["B16"] = "=B14-B15"
t["C16"] = "=C14-C15"
for r in (14, 15, 16):
    t.cell(r, 4, f"=C{r}-B{r}")
    t.cell(r, 5, f"=IF(B{r}=0,0,D{r}/B{r})")
    for col in range(1, 6):
        t.cell(r, col).font = Font(name="Calibri", bold=True, size=11)
        t.cell(r, col).fill = total_fill
        t.cell(r, col).border = thin
    chf(t.cell(r, 2))
    chf(t.cell(r, 3))
    chf(t.cell(r, 4))
    pct(t.cell(r, 5))

for i, w in enumerate([36, 14, 14, 12, 10], 1):
    t.column_dimensions[get_column_letter(i)].width = w

# --- Corrigé formateur ---
c = wb.create_sheet("Corrige_formateur")
c["A1"] = "Corrigé indicatif — Mars (pour le formateur)"
c["A1"].font = title_font
c["A3"] = "Synthèse attendue (idée, pas texte imposé)"
c["A3"].font = section_font
c["A4"] = (
    "Mars est nettement sous budget : le résultat souffre surtout d'un CA en retrait "
    "(−18'000 CHF), amplifié par plus de freelances (+4'000). Les autres écarts sont "
    "secondaires. À confirmer avec le métier : reports de missions et couverture freelances."
)
c["A4"].alignment = Alignment(wrap_text=True)
c.merge_cells("A4:B6")

c["A8"] = "3 postes qui devraient ressortir"
c["A8"].font = section_font
c["A9"] = "1. CA honoraires — défavorable −18'000 (−18,0 %)"
c["A10"] = "2. Sous-traitance — défavorable +4'000 (+50 %)"
c["A11"] = (
    "3. Marketing — défavorable +1'200 "
    "(ou personnel +1'500 — les deux sont acceptables en 3e)"
)
c["A13"] = "Questions métier possibles"
c["A13"].font = section_font
c["A14"] = "• Les 2 missions reportées sont-elles rebookées en avril ?"
c["A15"] = "• Le freelance remplace-t-il une absence ponctuelle ou un sous-effectif structurel ?"
c["A16"] = "• La campagne LinkedIn a-t-elle un ROI mesurable ce trimestre ?"
c["A18"] = "Prochaine action type"
c["A18"].font = section_font
c["A19"] = (
    "Mettre à jour le forecast avril avec les missions reportées + plafond freelance."
)
c["A21"] = "Ne montrez pas cette feuille aux participants avant la fin de l'exercice."
c["A21"].fill = note_fill
c["A21"].font = Font(name="Calibri", bold=True, size=11)
c.column_dimensions["A"].width = 95

out = Path(__file__).resolve().parent / "Budget_vs_Realise_AlpesConseil.xlsx"
wb.save(out)
print(out)
